from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from app.config import Principal, Settings
from app.main import create_app


def test_upload_retrieve_cite_feedback_and_tenant_isolation(tmp_path: Path):
    settings = Settings(
        app_env="test",
        data_dir=tmp_path,
        database_url=f"sqlite:///{tmp_path / 'e2e.db'}",
        qdrant_url=None,
        qdrant_path=tmp_path / "qdrant",
        embedding_model="BAAI/bge-small-en-v1.5",
        max_upload_bytes=500_000,
        chunk_words=60,
        chunk_overlap_words=10,
        api_keys={
            "admin-a": Principal("tenant-a", "admin"),
            "viewer-a": Principal("tenant-a", "viewer"),
            "viewer-b": Principal("tenant-b", "viewer"),
        },
    )
    app = create_app(settings)
    doc = DocxDocument()
    doc.add_heading("Lockout process", level=1)
    doc.add_paragraph(
        "Before maintenance, isolate hazardous energy sources and apply locks and tags. "
        "Verify that equipment cannot restart before work begins. "
        "This controlled sequence protects workers during servicing."
    )
    stream = BytesIO()
    doc.save(stream)
    content = stream.getvalue()
    with TestClient(app) as client:
        up = client.post(
            "/api/v1/documents/upload",
            headers={"X-API-Key": "admin-a"},
            files={
                "file": (
                    "lockout.docx",
                    content,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert up.status_code == 201, up.text
        assert up.json()["status"] == "indexed"
        duplicate = client.post(
            "/api/v1/documents/upload",
            headers={"X-API-Key": "admin-a"},
            files={
                "file": (
                    "lockout.docx",
                    content,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert duplicate.status_code == 201
        assert duplicate.json()["status"] == "duplicate"
        reindex = client.post("/api/v1/admin/reindex", headers={"X-API-Key": "admin-a"})
        assert reindex.status_code == 200
        assert reindex.json()["documents_reindexed"] == 1
        assert reindex.json()["chunks_reembedded"] > 0
        ask = client.post(
            "/api/v1/ask",
            headers={"X-API-Key": "viewer-a"},
            json={"query": "What should workers do before maintenance begins?"},
        )
        assert ask.status_code == 200, ask.text
        body = ask.json()
        assert not body["abstained"]
        assert body["citations"] and body["citations"][0]["title"] == "lockout.docx"
        assert "isolate hazardous energy" in body["answer"].lower()
        feedback = client.post(
            "/api/v1/feedback",
            headers={"X-API-Key": "viewer-a"},
            json={
                "conversation_id": body["conversation_id"],
                "rating": 5,
                "citation_correct": True,
            },
        )
        assert feedback.status_code == 201
        other = client.post(
            "/api/v1/search",
            headers={"X-API-Key": "viewer-b"},
            json={"query": "hazardous energy maintenance"},
        )
        assert other.status_code == 200
        assert other.json()["results"] == []
        assert client.get("/api/v1/documents", headers={"X-API-Key": "viewer-b"}).json() == []
        denied = client.post(
            "/api/v1/documents/upload",
            headers={"X-API-Key": "viewer-a"},
            files={"file": ("lockout.docx", content)},
        )
        assert denied.status_code == 403
