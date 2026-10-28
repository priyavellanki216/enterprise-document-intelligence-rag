from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader


@dataclass(frozen=True)
class PageText:
    page: int | None
    text: str


def parse_document(filename: str, content: bytes, max_bytes: int) -> list[PageText]:
    if not content:
        raise ValueError("Uploaded document is empty")
    if len(content) > max_bytes:
        raise ValueError(f"File exceeds the {max_bytes}-byte upload limit")
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        try:
            reader = PdfReader(BytesIO(content), strict=True)
            if reader.is_encrypted:
                raise ValueError("Encrypted PDFs are not supported")
            pages = [
                PageText(i, (page.extract_text() or "").strip())
                for i, page in enumerate(reader.pages, 1)
            ]
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f"Could not parse PDF: {exc}") from exc
        if not pages or not any(page.text for page in pages):
            raise ValueError("PDF contains no extractable text; scanned-PDF OCR is not enabled")
        return pages
    if suffix == ".docx":
        try:
            doc = DocxDocument(BytesIO(content))
        except Exception as exc:
            raise ValueError(f"Could not parse DOCX: {exc}") from exc
        pieces: list[str] = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                pieces.append(text)
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    pieces.append(" | ".join(cells))
        if not pieces:
            raise ValueError("DOCX contains no extractable text")
        return [PageText(None, "\n".join(pieces))]
    raise ValueError("Unsupported file type; upload PDF or DOCX")


def content_hash(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def chunk_pages(pages: list[PageText], words_per_chunk: int, overlap_words: int) -> list[dict]:
    chunks: list[dict] = []
    step = words_per_chunk - overlap_words
    heading = ""
    for page in pages:
        lines = [re.sub(r"\s+", " ", x).strip() for x in page.text.splitlines() if x.strip()]
        buffer: list[str] = []
        for line in lines:
            is_heading = len(line) <= 130 and (
                line.isupper() or bool(re.match(r"^(\d+(?:\.\d+)*\s+|[A-Z][.)]\s+)", line))
            )
            if is_heading and len(line.split()) < 18:
                heading = line
            buffer.extend(line.split())
        for start in range(0, len(buffer), step):
            text = " ".join(buffer[start : start + words_per_chunk]).strip()
            if len(text) >= 40:
                chunks.append({"page": page.page, "section": heading, "text": text})
            if start + words_per_chunk >= len(buffer):
                break
    return chunks
